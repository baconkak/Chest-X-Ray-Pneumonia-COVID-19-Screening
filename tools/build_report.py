"""Format the complete Thai text report as an editable Word document.

Run with the Python executable returned by load_workspace_dependencies.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'report_th.txt'
OUTPUT = ROOT / 'report_th.docx'
FONT = 'Leelawadee UI'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
section.footer_distance = Inches(.35)


def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:'+key), FONT)
    fonts.attrib.pop(qn('w:asciiTheme'), None)
    fonts.attrib.pop(qn('w:hAnsiTheme'), None)
    for tag in ['sz', 'szCs']:
        el = OxmlElement('w:'+tag); el.set(qn('w:val'), str(int(size*2))); rpr.append(el)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'th-TH'); lang.set(qn('w:bidi'), 'th-TH'); rpr.append(lang)
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for italic in list(rpr.iter(qn('w:iCs'))):
        italic.set(qn('w:val'), '0')


for name, size, bold in [('Normal',12,False), ('Title',22,True), ('Subtitle',13,False),
                         ('Heading 1',16,True), ('Heading 2',13,True),
                         ('List Bullet',12,False), ('List Number',12,False), ('Caption',10.5,False)]:
    set_font(doc.styles[name], size, bold)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.18
normal.space_after = Pt(7)
normal.widow_control = True
for name in ['Heading 1','Heading 2']:
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(16); pf.space_after = Pt(7); pf.keep_with_next=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)

# Page numbers are useful for this long academic report.
p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); r._r.append(field)
doc.core_properties.title='รายงานโครงงานการคัดกรอง Pneumonia จากภาพเอกซเรย์ทรวงอก'
doc.core_properties.subject='Deep Learning และการวิเคราะห์ขอบเขต COVID 19'
doc.core_properties.author=''
doc.core_properties.keywords='Pneumonia NIH ChestXray14 MLP CNN ResNet18'


def clean_heading(s):
    return re.sub(r'[^\w\s\u0E00-\u0E7F]', ' ', s).strip()


def math_run(text):
    r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; r.append(t); return r


def fraction(a,b):
    f=OxmlElement('m:f')
    for tag, text in [('num',a),('den',b)]:
        e=OxmlElement('m:'+tag); e.append(math_run(text)); f.append(e)
    return f


def equation(text, parts=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(9)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    mp=OxmlElement('m:oMathPara'); m=OxmlElement('m:oMath')
    for item in parts or [text]:
        m.append(math_run(item) if isinstance(item,str) else item)
    mp.append(m); p._p.append(mp)


def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,row_values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j,(cell,text) in enumerate(zip(row.cells,row_values)):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margin.append(e)
            tcpr.append(margin)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F5F7FA' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.1
            if j>0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(text); r.font.size=Pt(10.5); r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


source=SOURCE.read_text(encoding='utf-8')
lines=source.splitlines()
doc.add_paragraph('รายงานโครงงาน Deep Learning', 'Title')
doc.add_paragraph('การคัดกรอง Pneumonia จากภาพเอกซเรย์ทรวงอก\nและการวิเคราะห์ขอบเขต COVID 19', 'Subtitle')
for line in lines[3:7]: doc.add_paragraph(line)

# Preserve source content and paragraph boundaries while removing text-only rulers.
i=8
while i<len(lines):
    line=lines[i].strip()
    if not line or re.fullmatch(r'[-=]+',line): i+=1; continue
    if i+1<len(lines) and re.fullmatch(r'-+',lines[i+1].strip()):
        doc.add_paragraph(clean_heading(line),'Heading 1'); i+=2; continue
    if re.match(r'^6\.[1-4]\s',line):
        doc.add_paragraph(clean_heading(line),'Heading 2'); i+=1; continue
    if line.startswith('ตาราง A:'):
        p=doc.add_paragraph('ตาราง A ผล Validation สำหรับเลือกโมเดล','Heading 2')
        table(['โมเดล','AP','AUROC','Threshold','Recall','Precision'],
              [[n]+['รอผลจริง']*5 for n in ['Majority','MLP','SmallCNN','ResNet18']],
              [1.1,1.05,1.05,1.3,1.2,1.2])
        doc.add_paragraph('เติมจาก validation_comparison.csv โดย majority validation baseline ต้องประเมินแยกหากต้องการรายงาน')
        i+=1
        while i<len(lines) and not lines[i].startswith('Selected model'): i+=1
        continue
    if line.startswith('ตาราง B:'):
        doc.add_paragraph('ตาราง B ผล Final held out test','Heading 2')
        table(['โมเดล','AP และ\n95% group CI','AUROC','Recall','Precision','Specificity','F1'],
              [[n]+['รอผลจริง']*6 for n in ['Majority','MLP','SmallCNN','ResNet18']],
              [1.0,1.35,.85,.9,.95,1.0,.85])
        doc.add_paragraph('เติมจาก final_test/test_comparison.csv ห้ามใช้ผล test เลือกโมเดลใหม่')
        i+=1
        while i<len(lines) and not lines[i].startswith('แสดง TN'): i+=1
        continue
    if line.startswith('ตาราง C:'):
        doc.add_paragraph('ตาราง C ต้นทุนทรัพยากร','Heading 2')
        table(['โมเดล','Parameters','เวลาฝึกรวม\nวินาที','Best epoch','GPU และรุ่น'],
              [[n]+['รอผลจริง']*4 for n in ['MLP','SmallCNN','ResNet18']],
              [1.1,1.5,1.5,1.3,1.5])
        doc.add_paragraph('เติมจาก run.json และ history.csv ของแต่ละโมเดล')
        i+=1
        while i<len(lines) and lines[i].strip(): i+=1
        continue
    if line.startswith('L(z,y) ='):
        equation('L(z,y) = −[y log p + (1−y) log(1−p)]'); i+=1; continue
    if line.startswith('max(z,0) - zy'):
        equation('max(z,0) − zy + log(1 + exp(−|z|))'); i+=1; continue
    if line.startswith('w_pos ='):
        equation('', ['w_pos = ',fraction('N_negative_train','N_positive_train')]); i+=1; continue
    if line.startswith('L_weighted ='):
        equation('L_weighted = −[w_pos y log p + (1−y) log(1−p)]'); i+=1; continue
    if line.startswith('epoch_loss ='):
        equation('', ['epoch_loss = ',fraction('Σₖ batch_sizeₖ × batch_mean_lossₖ','Σₖ batch_sizeₖ')]); i+=1; continue
    if line.startswith('AP = sum_k'):
        equation('AP = Σₖ (Recallₖ − Recallₖ₋₁) × Precisionₖ'); i+=1; continue
    if line.startswith('F1 ='):
        equation('', ['F1 = ',fraction('2TP','2TP + FP + FN')]); i+=1; continue
    if line.startswith('Brier ='):
        equation('Brier = mean((score − y)²)')
        doc.add_paragraph(line.split(':',1)[1].strip()); i+=1; continue
    if re.match(r'^(Recall|Precision|Specificity|Accuracy|F1) =',line):
        match=re.match(r'^(\w+) = ([^:]+)\s*:\s*(.+)$',line)
        if match:
            metric,formula,meaning=match.groups()
            a,b=formula.strip().split('/',1)
            equation('', [metric+' = ',fraction(a.strip('()'),b.strip('()'))])
            doc.add_paragraph(meaning); i+=1; continue
    if line.startswith('- '):
        # Join continuation lines of one bullet without turning every source wrap into a paragraph.
        block=[line[2:]]; i+=1
        while i<len(lines) and lines[i].startswith('    ') and not lines[i].lstrip().startswith('- '):
            block.append(lines[i].strip()); i+=1
        doc.add_paragraph(' '.join(block),'List Bullet'); continue
    if re.match(r'^\(\d+\)\s',line):
        p=doc.add_paragraph(line); p.paragraph_format.left_indent=Inches(.18); i+=1; continue
    if re.match(r'^\[\d+\]',line):
        block=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^\[\d+\]',lines[i].strip()):
            block.append(lines[i].strip()); i+=1
        p=doc.add_paragraph(' '.join(block)); p.paragraph_format.space_after=Pt(10)
        for r in p.runs: r.font.size=Pt(10.5)
        continue
    block=[line]; i+=1
    while i<len(lines) and lines[i].strip():
        nxt=lines[i].strip()
        if re.fullmatch(r'[-=]+',nxt) or re.match(r'^6\.[1-4]\s',nxt): break
        if nxt.startswith(('- ','ตาราง ')) or re.match(r'^\(\d+\)\s|^\[\d+\]',nxt): break
        if i+1<len(lines) and re.fullmatch(r'-+',lines[i+1].strip()): break
        # Preserve mathematical display lines for the next loop iteration.
        if nxt.startswith(('L(z,y) =','max(z,0) - zy','w_pos =','L_weighted =','epoch_loss =','AP = sum_k')): break
        if re.match(r'^(Recall|Precision|Specificity|Accuracy|F1) =',nxt): break
        block.append(nxt); i+=1
    text=' '.join(block)
    p=doc.add_paragraph(text)
    if len(text)<100 and (text.endswith(':') or text.startswith(('H1:','H2:','H3:'))):
        p.runs[0].bold=True; p.paragraph_format.keep_with_next=True

doc.save(OUTPUT)
print(OUTPUT)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
